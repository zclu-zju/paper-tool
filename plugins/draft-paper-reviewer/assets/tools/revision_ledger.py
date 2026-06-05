#!/usr/bin/env python3
"""Build the draft-paper-reviewer revision ledger workbook from JSONL records."""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover - exercised in target repos
    raise SystemExit(
        "Missing dependency: openpyxl. Install with:\n"
        "python3 -m pip install -r .codex/tools/draft-paper-reviewer/requirements.txt"
    ) from exc


FIELDS = [
    "scope_id",
    "scope_name",
    "round",
    "reviewer_agent",
    "reviser_agent",
    "review_record",
    "review_score",
    "review_status",
    "change_record",
    "changed_files",
    "evidence_ids",
    "convention_ids",
    "retention_gate_ids",
    "symbol_ids",
    "revision_task_ids",
    "acceptance_criteria",
    "next_action",
    "timestamp_utc",
]

REVIEW_COLUMNS = [
    "round",
    "reviewer_agent",
    "review_score",
    "review_status",
    "review_record",
    "acceptance_criteria",
    "evidence_ids",
    "convention_ids",
    "retention_gate_ids",
    "symbol_ids",
]

CHANGE_COLUMNS = [
    "round",
    "reviser_agent",
    "change_record",
    "changed_files",
    "revision_task_ids",
    "convention_ids",
    "retention_gate_ids",
    "symbol_ids",
    "next_action",
    "timestamp_utc",
]


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        raise SystemExit(f"Ledger JSONL not found: {path}")
    rows: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as handle:
        for line_no, line in enumerate(handle, 1):
            line = line.strip()
            if not line:
                continue
            try:
                item = json.loads(line)
            except json.JSONDecodeError as exc:
                raise SystemExit(f"Invalid JSON on {path}:{line_no}: {exc}") from exc
            rows.append({field: item.get(field, "") for field in FIELDS})
    return rows


def sheet_name(scope_id: str, used: set[str]) -> str:
    base = re.sub(r"[:\\/?*\[\]]", "_", scope_id or "unknown_scope")[:31] or "unknown_scope"
    name = base
    suffix = 1
    while name in used:
        token = f"_{suffix}"
        name = f"{base[:31 - len(token)]}{token}"
        suffix += 1
    used.add(name)
    return name


def autosize(ws) -> None:
    for column in ws.columns:
        max_len = 0
        letter = get_column_letter(column[0].column)
        for cell in column:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(value), 80))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.column_dimensions[letter].width = max(12, min(max_len + 2, 72))


def style_header(ws, row: int = 1) -> None:
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[row]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_workbook(rows: list[dict[str, Any]], xlsx_path: Path) -> None:
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(
        [
            "scope_id",
            "scope_name",
            "rounds",
            "latest_score",
            "latest_status",
            "latest_next_action",
            "latest_timestamp_utc",
        ]
    )

    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[str(row.get("scope_id") or "unknown_scope")].append(row)

    for scope_id, items in sorted(grouped.items()):
        items.sort(key=lambda item: str(item.get("round", "")))
        latest = items[-1]
        summary.append(
            [
                scope_id,
                latest.get("scope_name", ""),
                len(items),
                latest.get("review_score", ""),
                latest.get("review_status", ""),
                latest.get("next_action", ""),
                latest.get("timestamp_utc", ""),
            ]
        )

    style_header(summary)
    autosize(summary)

    used = {"Summary"}
    for scope_id, items in sorted(grouped.items()):
        ws = wb.create_sheet(sheet_name(scope_id, used))
        ws.append(["Review Record", "", "", "", "", "", "", "Change Record", "", "", "", "", "", ""])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(REVIEW_COLUMNS))
        ws.merge_cells(
            start_row=1,
            start_column=len(REVIEW_COLUMNS) + 1,
            end_row=1,
            end_column=len(REVIEW_COLUMNS) + len(CHANGE_COLUMNS),
        )
        ws.cell(row=1, column=1).font = Font(bold=True)
        ws.cell(row=1, column=len(REVIEW_COLUMNS) + 1).font = Font(bold=True)
        ws.append(REVIEW_COLUMNS + CHANGE_COLUMNS)
        style_header(ws, 2)
        for item in items:
            ws.append([item.get(col, "") for col in REVIEW_COLUMNS + CHANGE_COLUMNS])
        ws.freeze_panes = "A3"
        autosize(ws)

    wb.save(xlsx_path)


def write_csv_exports(rows: list[dict[str, Any]], csv_dir: Path) -> None:
    csv_dir.mkdir(parents=True, exist_ok=True)
    with (csv_dir / "all_records.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(rows)

    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[str(row.get("scope_id") or "unknown_scope")].append(row)
    for scope_id, items in grouped.items():
        safe = re.sub(r"[^A-Za-z0-9_.-]+", "_", scope_id) or "unknown_scope"
        with (csv_dir / f"{safe}.csv").open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=FIELDS)
            writer.writeheader()
            writer.writerows(items)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--jsonl", required=True, type=Path)
    parser.add_argument("--xlsx", required=True, type=Path)
    parser.add_argument("--csv-dir", required=True, type=Path)
    args = parser.parse_args()

    rows = read_jsonl(args.jsonl)
    write_workbook(rows, args.xlsx)
    write_csv_exports(rows, args.csv_dir)
    print(f"Wrote XLSX ledger: {args.xlsx}")
    print(f"Wrote CSV exports: {args.csv_dir}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
